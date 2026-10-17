#!/usr/bin/env python3
"""
Coding sheets for the validation sample of split clusters (final pre-submission audit, GIQ v25).

build (default) writes, from reports/split_audit_sample.csv (contracts) and reports/split_audit_codes.csv (codes):
  reports/validation/split_audit_verify.xlsx   all 200 clusters with the proposed code, for the author to verify
                                                (column "verified code"); summary sheet counts codes by period
  reports/validation/second_coder_blind.xlsx   50 clusters (25 pre-reform, 25 post-Act, random), shuffled, with
                                                period, dates and proposed codes hidden, for an independent coder
  reports/validation/second_coder_key.csv      blind id -> cluster, period (keep away from the second coder)
kappa reads the filled sheets and reports agreement and Cohen's kappa:
  reports/validation/split_audit_agreement.md

Codes: A = one purchase divided (same item or work in lots, sections, sites or periods);
       B = separate needs (different items or works); C = cannot tell from the names.
Vendor names are not in the sample (personal data under the PDPA).

Usage (Anaconda Prompt, env egp-xai):  python src\\split_audit_sheets.py [build|kappa]
"""
import re
import sys
from pathlib import Path

import numpy as np
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"
VAL = REPORTS / "validation"
SEED, N_BLIND = 2026, 25
CODES = ("A", "B", "C")
RULES = [
    ("A", "One purchase divided: the same item or the same type of work bought in two or more contracts "
          "(e.g., concrete roads in several villages, diesel in several orders, one drug in several orders)."),
    ("B", "Separate needs: different items or works only, which happen to share a vendor and a signing day."),
    ("C", "Cannot tell: project names too generic to decide."),
]


def load():
    s = pd.read_csv(REPORTS / "split_audit_sample.csv", dtype={"project_id": str})
    c = pd.read_csv(REPORTS / "split_audit_codes.csv")
    assert set(s.cluster) == set(c.cluster), "sample and codes do not cover the same clusters"
    return s, c


def mask_period(name):
    """Hide cues to the period in a project name: the method (ตกลงราคา before the 2017 Act, เฉพาะเจาะจง after) and
    Buddhist-era years."""
    name = re.sub(r"\s*\(?(โดย|ด้วย)?\s*วิธี[^\s\[\)]*\)?", "", name)
    return re.sub(r"(?<!\d)25[5-6]\d(?!\d)", "25xx", name)


def cluster_table(s, blind=False):
    if blind:
        s = s.assign(project_name=s.project_name.fillna("").map(mask_period))
    s = s.sort_values(["cluster", "agreed_price"], ascending=[True, False])
    g = s.groupby("cluster")
    return pd.DataFrame({
        "agency": g.agency.first(),
        "contracts": g.size(),
        "total_THB": g.agreed_price.sum().round(0),
        "project_names": g.apply(lambda d: "\n".join(f"{i + 1}. {n} [{p:,.0f}]" for i, (n, p) in
                                                     enumerate(zip(d.project_name, d.agreed_price)))),
    }).reset_index()


def build():
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    VAL.mkdir(parents=True, exist_ok=True)
    s, c = load()
    T = cluster_table(s).merge(c, on="cluster")
    F = Font(name="Arial", size=10)
    B = Font(name="Arial", size=10, bold=True)
    INPUT = PatternFill("solid", fgColor="FFFF00")
    wrap = Alignment(wrap_text=True, vertical="top")

    def instructions(ws, title, lines):
        ws["A1"] = title
        ws["A1"].font = Font(name="Arial", size=12, bold=True)
        r = 3
        for t in lines:
            ws.cell(r, 1, t).font = F
            ws.cell(r, 1).alignment = Alignment(wrap_text=True)
            r += 1
        r += 1
        ws.cell(r, 1, "Code").font = B
        ws.cell(r, 2, "Meaning").font = B
        for k, t in RULES:
            r += 1
            ws.cell(r, 1, k).font = F
            ws.cell(r, 2, t).font = F
            ws.cell(r, 2).alignment = Alignment(wrap_text=True)
        ws.column_dimensions["A"].width = 100 if len(lines[0]) > 60 else 12
        ws.column_dimensions["B"].width = 90

    def sheet(ws, df, cols, widths, inputs):
        for j, (h, w) in enumerate(zip(cols, widths), 1):
            ws.cell(1, j, h).font = B
            ws.column_dimensions[ws.cell(1, j).column_letter].width = w
        for i, row in enumerate(df.itertuples(index=False), 2):
            for j, v in enumerate(row, 1):
                cell = ws.cell(i, j, v)
                cell.font = F
                cell.alignment = wrap
            for j in inputs:
                ws.cell(i, j).fill = INPUT
        ws.freeze_panes = "B2"
        dv = DataValidation(type="list", formula1='"A,B,C"', allow_blank=True)
        ws.add_data_validation(dv)
        for j in inputs[:1]:
            L = ws.cell(1, j).column_letter
            dv.add(f"{L}2:{L}{len(df) + 1}")

    # 1. verification workbook (all 200 clusters)
    wb = Workbook()
    ws = wb.active
    ws.title = "Instructions"
    instructions(ws, "Verification of the split-cluster validation codes (Section 6 of the manuscript)", [
        "Each row of sheet Clusters is one same-day split cluster: one buying unit, one vendor, one signing day, two or "
        "more contracts of at most THB 500,000 summing to THB 0.5–2 million. Vendor names are withheld (PDPA).",
        "Column 'proposed code' holds the code assigned in preparing the manuscript. Read the project names and enter "
        "your own code in the yellow column 'verified code' (A, B or C); add a note when you disagree.",
        "Example: cluster 2558-001 lists three orders of concrete poles of different sizes from the same vendor on one "
        "day → A if you judge them one purchase divided, B if separate needs.",
        "Sheet Summary counts the codes by period from both columns and the share of rows where they agree; "
        "its figures update as you fill in the verified codes.",
    ])
    ws = wb.create_sheet("Clusters")
    V = T[["cluster", "period", "agency", "contracts", "total_THB", "project_names", "code"]].copy()
    V["verified"] = None
    V["note"] = None
    sheet(ws, V, ["cluster", "period", "agency", "contracts", "total THB", "project names [agreed price, THB]",
                  "proposed code", "verified code", "note"], [10, 14, 24, 9, 12, 90, 9, 9, 30], [8, 9])
    ws["H1"].comment = Comment("Enter A, B or C. Leave blank until checked.", "split_audit_sheets.py")
    n = len(V) + 1
    ws = wb.create_sheet("Summary")
    ws["A1"], ws["B1"], ws["C1"], ws["D1"], ws["E1"] = "period", "code", "proposed", "verified", "verified share"
    for cell in ws[1]:
        cell.font = B
    r = 2
    for per in ("pre (FY2015)", "post-Act"):
        for k in CODES:
            ws.cell(r, 1, per); ws.cell(r, 2, k)
            ws.cell(r, 3, f'=COUNTIFS(Clusters!$B$2:$B${n},$A{r},Clusters!$G$2:$G${n},$B{r})')
            ws.cell(r, 4, f'=COUNTIFS(Clusters!$B$2:$B${n},$A{r},Clusters!$H$2:$H${n},$B{r})')
            ws.cell(r, 5, f'=IFERROR(D{r}/SUMIFS($D$2:$D$7,$A$2:$A$7,$A{r}),"")')
            ws.cell(r, 5).number_format = "0.0%"
            r += 1
    ws.cell(r + 1, 1, "rows verified"); ws.cell(r + 1, 4, f'=COUNTA(Clusters!$H$2:$H${n})')
    ws.cell(r + 2, 1, "rows where verified = proposed")
    ws.cell(r + 2, 4, f'=SUMPRODUCT((Clusters!$H$2:$H${n}<>"")*(Clusters!$H$2:$H${n}=Clusters!$G$2:$G${n}))')
    ws.cell(r + 3, 1, "agreement"); ws.cell(r + 3, 4, f'=IFERROR(D{r + 2}/D{r + 1},"")')
    ws.cell(r + 3, 4).number_format = "0.0%"
    for row in ws.iter_rows():
        for cell in row:
            cell.font = B if cell.row == 1 else F
    for L, w in zip("ABCDE", (32, 8, 10, 10, 14)):
        ws.column_dimensions[L].width = w
    wb.save(VAL / "split_audit_verify.xlsx")

    # 2. blind second-coder workbook (25 pre, 25 post, shuffled; period, dates and codes hidden)
    rng = np.random.default_rng(SEED)
    pick = pd.concat([T[T.period == p].iloc[rng.choice((T.period == p).sum(), N_BLIND, replace=False)]
                      for p in ("pre (FY2015)", "post-Act")])
    pick = pick.iloc[rng.permutation(len(pick))].reset_index(drop=True)
    pick = pick[["cluster", "period"]].merge(cluster_table(s, blind=True), on="cluster")    # period cues masked
    pick["blind_id"] = [f"X{i + 1:02d}" for i in range(len(pick))]
    pick[["blind_id", "cluster", "period"]].to_csv(VAL / "second_coder_key.csv", index=False, encoding="utf-8")
    wb = Workbook()
    ws = wb.active
    ws.title = "Instructions"
    instructions(ws, "Coding of 50 split clusters (independent second coder)", [
        "Each row of sheet Clusters is one group of contracts that one buying unit signed with one vendor on one day. "
        "Every contract is at most THB 500,000 and the group sums to THB 0.5–2 million.",
        "Method names and years in the project names are masked (25xx), so that the period cannot be inferred. "
        "Read the project names and enter one code in the yellow column 'code' (A, B or C). Code each row on its own; "
        "do not consult the manuscript, the other coder or any other file of the project while coding.",
        "Example: three orders of concrete poles of different sizes from one vendor on one day. Decide whether they "
        "are one purchase divided (A) or separate needs (B); if the names do not tell, use C.",
    ])
    ws = wb.create_sheet("Clusters")
    X = pick[["blind_id", "agency", "contracts", "total_THB", "project_names"]].copy()
    X["code"] = None
    X["note"] = None
    sheet(ws, X, ["id", "agency", "contracts", "total THB", "project names [agreed price, THB]", "code", "note"],
          [7, 24, 9, 12, 90, 8, 30], [6, 7])
    wb.save(VAL / "second_coder_blind.xlsx")
    print("→", VAL / "split_audit_verify.xlsx", "|", VAL / "second_coder_blind.xlsx", "|", VAL / "second_coder_key.csv")


def kappa(a, b):
    labs = sorted(set(a) | set(b))
    M = pd.crosstab(pd.Categorical(a, labs), pd.Categorical(b, labs), dropna=False).values.astype(float)
    n = M.sum()
    po = np.trace(M) / n
    pe = (M.sum(0) * M.sum(1)).sum() / n ** 2
    return po, (po - pe) / (1 - pe) if pe < 1 else np.nan, M, labs


def agreement():
    _, c = load()
    v = pd.read_excel(VAL / "split_audit_verify.xlsx", sheet_name="Clusters")
    key = pd.read_csv(VAL / "second_coder_key.csv")
    b = pd.read_excel(VAL / "second_coder_blind.xlsx", sheet_name="Clusters").rename(columns={"id": "blind_id"})
    L = ["# Validation codes: verification and inter-coder agreement", ""]
    vv = v.dropna(subset=["verified code"])
    L.append(f"Verified by the author: {len(vv)} of {len(v)} clusters; verified = proposed in "
             f"{(vv['verified code'] == vv['proposed code']).mean():.1%}." if len(vv) else "Not yet verified.")
    final = v.assign(final=v["verified code"].fillna(v["proposed code"]))[["cluster", "period", "final"]]
    L += ["", "| period | A | B | C | n |", "|---|---|---|---|---|"]
    for p, d in final.groupby("period", sort=False):
        L.append(f"| {p} | " + " | ".join(f"{(d.final == k).mean():.0%}" for k in CODES) + f" | {len(d)} |")
    m = key.merge(b[["blind_id", "code"]], on="blind_id").merge(final, on=["cluster", "period"]).dropna(subset=["code"])
    if len(m):
        po, k, M, labs = kappa(m.final.str.strip(), m.code.astype(str).str.strip().str.upper())
        L += ["", f"Second coder: {len(m)} clusters coded. Agreement {po:.1%}; Cohen's kappa {k:.2f}.", "",
              "| first \\ second | " + " | ".join(labs) + " |", "|---" * (len(labs) + 1) + "|"]
        L += [f"| {r} | " + " | ".join(f"{x:.0f}" for x in M[i]) + " |" for i, r in enumerate(labs)]
    else:
        L += ["", "Second coder: no codes entered yet."]
    (VAL / "split_audit_agreement.md").write_text("\n".join(L) + "\n", encoding="utf-8")
    print("\n".join(L))


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    (agreement if sys.argv[1:] == ["kappa"] else build)()
